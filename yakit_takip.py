import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
from datetime import datetime, timedelta
import sqlite3
import os
import sys
import matplotlib
matplotlib.use('TkAgg')
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import matplotlib.pyplot as plt
from tkcalendar import DateEntry
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

class AraçTakipUygulaması:
    def __init__(self, root):
        self.root = root
        self.setup_database()
        self.setup_styles()   
        self.setup_ui()
        self.load_initial_data()
        self.check_notifications()
        
        self.transaction_type = tk.StringVar(value="IN") 
        
    def setup_database(self):
        """Veritabanı tablolarını oluşturur veya bağlantı kurar"""
        self.db_name = "arac_takip.db"
        
        if getattr(sys, 'frozen', False):
            application_path = os.path.dirname(sys.executable)
        else:
            application_path = os.path.dirname(os.path.abspath(__file__))
        
        self.db_path = os.path.join(application_path, self.db_name)
        
        try:
            self.conn = sqlite3.connect(self.db_path)
            self.conn.execute("PRAGMA foreign_keys = ON")
            self.cursor = self.conn.cursor()
            
            # Tabloları oluştur (IF NOT EXISTS ile)
            self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS vehicles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                plate TEXT UNIQUE NOT NULL,
                model TEXT NOT NULL,
                km INTEGER DEFAULT 0,
                driver TEXT
            )""")
            
            self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS fuel_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                vehicle_id INTEGER NOT NULL,
                date TEXT NOT NULL,  -- ISO format: YYYY-MM-DD
                km INTEGER NOT NULL,
                amount REAL NOT NULL,
                price REAL NOT NULL,
                total REAL NOT NULL,
                FOREIGN KEY (vehicle_id) REFERENCES vehicles(id)
            )""")
            
            self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS fuel_tank (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                date TEXT NOT NULL,  -- ISO format: YYYY-MM-DD
                amount REAL NOT NULL,
                price REAL NOT NULL,
                total REAL NOT NULL,
                transaction_type TEXT NOT NULL CHECK(transaction_type IN ('IN', 'OUT')),
                notes TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )""")
            
            self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS maintenance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                vehicle_id INTEGER NOT NULL,
                date TEXT NOT NULL,  -- ISO format: YYYY-MM-DD
                km INTEGER NOT NULL,
                fault TEXT,
                repair TEXT,
                labor_cost REAL DEFAULT 0,
                material_cost REAL DEFAULT 0,
                FOREIGN KEY (vehicle_id) REFERENCES vehicles(id)
            )""")
            
            self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS inspections (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                vehicle_id INTEGER NOT NULL,
                date TEXT NOT NULL,  -- ISO format: YYYY-MM-DD
                km INTEGER NOT NULL,
                next_inspection_date TEXT,  -- ISO format: YYYY-MM-DD
                next_maintenance_date TEXT,  -- ISO format: YYYY-MM-DD
                next_maintenance_km INTEGER,
                FOREIGN KEY (vehicle_id) REFERENCES vehicles(id)
            )""")
            
            self.conn.commit()
            
            # Yakıt fiyatı ve depo durumu için değişkenler
            self.current_fuel_price = 20.0
            self.current_fuel_level = 0.0
            
            try:
                self.cursor.execute("SELECT price FROM fuel_tank WHERE transaction_type='IN' ORDER BY date DESC LIMIT 1")
                result = self.cursor.fetchone()
                if result:
                    self.current_fuel_price = float(result[0])
            except sqlite3.OperationalError:
                self.current_fuel_price = 20.0
            
            try:
                self.cursor.execute("""
                SELECT SUM(CASE 
                    WHEN transaction_type='IN' THEN amount 
                    ELSE -amount 
                END) FROM fuel_tank
                """)
                result = self.cursor.fetchone()
                self.current_fuel_level = float(result[0]) if result and result[0] else 0.0
            except sqlite3.OperationalError:
                self.current_fuel_level = 0.0
                
        except sqlite3.Error as e:
            messagebox.showerror("Veritabanı Hatası", f"Veritabanı bağlantısı kurulamadı: {str(e)}")
            self.root.destroy()
            raise

    def setup_styles(self):
        self.primary_color = "#2c3e50"
        self.secondary_color = "#3498db"
        self.accent_color = "#e74c3c"
        self.success_color = "#2ecc71"
        self.warning_color = "#f39c12"
        self.light_bg = "#ecf0f1"
        self.lighter_bg = "#f8f9fa"
        self.dark_text = "#2c3e50"
        self.light_text = "#ecf0f1"
        
        self.title_font = ('Segoe UI', 12, 'bold')
        self.subtitle_font = ('Segoe UI', 10, 'bold')
        self.normal_font = ('Segoe UI', 9)
        self.small_font = ('Segoe UI', 8)
        
        self.style = ttk.Style()
        self.style.theme_use('clam')
        
        self.style.configure('.', 
                           background=self.light_bg,
                           foreground=self.dark_text,
                           font=self.normal_font)
        
        self.style.configure('TFrame', background=self.light_bg)
        self.style.configure('Header.TFrame', background=self.primary_color)
        self.style.configure('Status.TFrame', background=self.primary_color)
        
        self.style.configure('TLabel', 
                           background=self.light_bg,
                           foreground=self.dark_text,
                           font=self.normal_font)
        self.style.configure('Title.TLabel', 
                           font=self.title_font,
                           foreground=self.primary_color)
        self.style.configure('Subtitle.TLabel', 
                           font=self.subtitle_font,
                           foreground=self.secondary_color)
        
        self.style.configure('TButton', 
                           font=self.subtitle_font,
                           borderwidth=1,
                           relief='raised',
                           padding=6)
        self.style.configure('Primary.TButton', 
                           foreground=self.light_text,
                           background=self.secondary_color,
                           borderwidth=0)
        self.style.map('Primary.TButton',
                      background=[('active', self.primary_color), ('pressed', self.accent_color)],
                      foreground=[('active', self.light_text), ('pressed', self.light_text)])
        
        self.style.configure('Danger.TButton', 
                           foreground=self.light_text,
                           background=self.accent_color,
                           borderwidth=0)
        self.style.map('Danger.TButton',
                      background=[('active', '#c0392b'), ('pressed', '#a93226')])
        
        self.style.configure('Success.TButton', 
                           foreground=self.light_text,
                           background=self.success_color,
                           borderwidth=0)
        self.style.map('Success.TButton',
                      background=[('active', '#27ae60'), ('pressed', '#219653')])
        
        self.style.configure('TEntry', 
                           fieldbackground="white",
                           foreground=self.dark_text,
                           insertcolor=self.dark_text,
                           padding=5,
                           bordercolor=self.secondary_color,
                           lightcolor=self.secondary_color)
        
        self.style.configure('TCombobox', 
                           fieldbackground="white",
                           foreground=self.dark_text,
                           selectbackground=self.secondary_color,
                           padding=5)
        
        self.style.configure('TNotebook', background=self.light_bg)
        self.style.configure('TNotebook.Tab', 
                           background=self.lighter_bg,
                           foreground=self.dark_text,
                           padding=[10, 5],
                           font=self.subtitle_font)
        self.style.map('TNotebook.Tab', 
                      background=[('selected', self.secondary_color)],
                      foreground=[('selected', self.light_text)])
        
        self.style.configure('Treeview', 
                           background="white",
                           foreground=self.dark_text,
                           rowheight=25,
                           fieldbackground="white",
                           font=self.normal_font,
                           bordercolor=self.light_bg,
                           lightcolor=self.light_bg)
        self.style.configure('Treeview.Heading', 
                           background=self.primary_color,
                           foreground=self.light_text,
                           font=self.subtitle_font,
                           padding=5)
        self.style.map('Treeview', 
                      background=[('selected', self.secondary_color)],
                      foreground=[('selected', self.light_text)])
        
        self.style.configure('Vertical.TScrollbar', 
                           background=self.light_bg,
                           arrowcolor=self.secondary_color,
                           troughcolor=self.light_bg)
        
        self.style.configure('TLabelframe', 
                           background=self.light_bg,
                           foreground=self.primary_color,
                           font=self.subtitle_font,
                           bordercolor=self.light_bg)
        self.style.configure('TLabelframe.Label', 
                           background=self.light_bg,
                           foreground=self.primary_color,
                           font=self.subtitle_font)
        
        self.style.configure('DateEntry', 
                           fieldbackground="white",
                           foreground=self.dark_text,
                           arrowcolor=self.secondary_color,
                           selectbackground=self.secondary_color)

    def setup_ui(self):
        try:
            self.root.title("Araç Takip Sistemi")
            self.root.geometry("1200x800")
            self.root.minsize(1000, 700)
            self.root.configure(bg=self.light_bg)
            
            # Header
            header_frame = ttk.Frame(self.root, style='Header.TFrame')
            header_frame.pack(fill=tk.X, padx=0, pady=0)
            
            ttk.Label(header_frame, 
                     text="ARAÇ TAKİP SİSTEMİ", 
                     style='Title.TLabel',
                     foreground=self.light_text,
                     background=self.primary_color).pack(pady=15)
            
            # Notebook (Sekmeler)
            self.notebook = ttk.Notebook(self.root)
            
            # Araç sekmesi
            self.create_arac_tab()
            
            # Yakıt sekmesi
            self.create_yakit_tab()
            
            # Depo sekmesi
            self.create_depo_tab()
            
            # Bakım sekmesi
            self.create_bakim_tab()
            
            # Muayene ve periyodik bakım sekmesi
            self.create_muayene_bakim_tab()
            
            # Rapor sekmesi
            self.create_rapor_tab()
            
            # Maliyet sekmesi
            self.create_maliyet_tab()
            
            self.notebook.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 5))
            
            # Durum çubuğu
            self.status_bar = ttk.Frame(self.root, style='Status.TFrame')
            self.fiyat_label = ttk.Label(
                self.status_bar, 
                text=f"Mevcut Yakıt Fiyatı: {self.current_fuel_price:.2f} TL",
                font=self.subtitle_font,
                foreground=self.light_text,
                background=self.primary_color
            )
            self.fiyat_label.pack(side=tk.RIGHT, padx=10, pady=5)
            self.status_bar.pack(fill=tk.X, padx=0, pady=0)
            
            # Menü oluştur
            self.create_menu()
            
        except Exception as e:
            messagebox.showerror("UI Hatası", f"Arayüz oluşturulurken hata: {str(e)}")
            self.root.destroy()
            raise

    def create_menu(self):
        menubar = tk.Menu(self.root, bg=self.light_bg, fg=self.dark_text, activebackground=self.secondary_color)
        
        # Dosya menüsü
        file_menu = tk.Menu(menubar, tearoff=0, bg=self.light_bg, fg=self.dark_text, activebackground=self.secondary_color)
        file_menu.add_command(label="Yakıt Fiyatı Güncelle", command=self.update_fuel_price)
        file_menu.add_command(label="Yedek Oluştur", command=self.create_backup)
        file_menu.add_command(label="Yedekten Geri Yükle", command=self.restore_backup)
        
        # Excel çıktı alt menüsü
        excel_menu = tk.Menu(file_menu, tearoff=0, bg=self.light_bg, fg=self.dark_text, activebackground=self.secondary_color)
        excel_menu.add_command(label="Yakıt İşlemleri", command=self.export_fuel_to_excel)
        excel_menu.add_command(label="Bakım İşlemleri", command=self.export_maintenance_to_excel)
        excel_menu.add_command(label="Maliyet Raporu", command=self.export_cost_report_to_excel)
        
        file_menu.add_cascade(label="Excel'e Aktar", menu=excel_menu)
        file_menu.add_separator()
        file_menu.add_command(label="Çıkış", command=self.on_closing)
        menubar.add_cascade(label="Dosya", menu=file_menu)

        # Uyarılar menüsü
        self.notification_menu = tk.Menu(menubar, tearoff=0, bg=self.light_bg, fg=self.dark_text, activebackground=self.secondary_color)
        self.notification_menu.add_command(label="Bakım Uyarıları", command=self.show_maintenance_notifications)
        self.notification_menu.add_command(label="Muayene Uyarıları", command=self.show_inspection_notifications)
        menubar.add_cascade(label="Uyarılar", menu=self.notification_menu)
        
        # Rapor menüsü
        report_menu = tk.Menu(menubar, tearoff=0, bg=self.light_bg, fg=self.dark_text, activebackground=self.secondary_color)
        report_menu.add_command(label="Maliyet Raporu", command=self.show_cost_report)
        report_menu.add_command(label="Tüketim Grafiği", command=self.show_consumption_graph)
        menubar.add_cascade(label="Raporlar", menu=report_menu)
        
        self.root.config(menu=menubar)

    def create_arac_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Araçlar", padding=5)
        
        # Add/edit vehicle panel
        add_frame = ttk.LabelFrame(frame, text="Araç Ekle/Düzenle", padding=10)
        add_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Label(add_frame, text="Plaka:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.plaka_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.plaka_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Model:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.model_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.model_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="KM:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.km_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.km_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Sürücü:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.driver_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.driver_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame = ttk.Frame(add_frame)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Kaydet", command=self.save_vehicle, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Temizle", command=self.clear_vehicle_form, style='TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Sil", command=self.delete_vehicle, style='Danger.TButton').pack(side=tk.LEFT, padx=5)
        
        # Vehicle list
        list_frame = ttk.LabelFrame(frame, text="Araç Listesi", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        columns = ("Plaka", "Model", "KM", "Sürücü")
        self.vehicle_tree = ttk.Treeview(list_frame, columns=columns, show="headings", selectmode="browse")
        
        for col in columns:
            self.vehicle_tree.heading(col, text=col, anchor="center")
            self.vehicle_tree.column(col, width=120, anchor="center")
        
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.vehicle_tree.yview)
        self.vehicle_tree.configure(yscrollcommand=scrollbar.set)
        
        self.vehicle_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.vehicle_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.vehicle_tree.tag_configure('evenrow', background="white")
        
        self.vehicle_tree.bind("<<TreeviewSelect>>", self.load_vehicle_data)

    def create_yakit_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Yakıt Kayıtları", padding=5)
        
        # Add fuel panel
        add_frame = ttk.LabelFrame(frame, text="Yakıt Ekle", padding=10)
        add_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Label(add_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.fuel_vehicle_combo = ttk.Combobox(add_frame, state="readonly", font=self.normal_font)
        self.fuel_vehicle_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="KM:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.fuel_km_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.fuel_km_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Miktar (L):", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.fuel_amount_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.fuel_amount_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Birim Fiyat:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.fuel_price_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.fuel_price_entry.insert(0, f"{self.current_fuel_price:.2f}")
        self.fuel_price_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Tarih:", style='Subtitle.TLabel').grid(row=4, column=0, sticky="w", pady=5, padx=5)
        self.fuel_date_entry = DateEntry(add_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.fuel_date_entry.grid(row=4, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame = ttk.Frame(add_frame)
        btn_frame.grid(row=5, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Kaydet", command=self.save_fuel_record, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Hesapla", command=self.calculate_fuel_cost, style='TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Sil", command=self.delete_fuel_record, style='Danger.TButton').pack(side=tk.LEFT, padx=5)
        
        # Fuel records list
        list_frame = ttk.LabelFrame(frame, text="Yakıt Kayıtları", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        columns = ("ID", "Tarih", "Plaka", "KM", "Miktar", "Birim Fiyat", "Toplam")
        self.fuel_tree = ttk.Treeview(list_frame, columns=columns, show="headings", selectmode="browse")
        
        for col in columns:
            self.fuel_tree.heading(col, text=col, anchor="center")
            self.fuel_tree.column(col, width=100, anchor="center")
        
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.fuel_tree.yview)
        self.fuel_tree.configure(yscrollcommand=scrollbar.set)
        
        self.fuel_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.fuel_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.fuel_tree.tag_configure('evenrow', background="white")
        
        self.fuel_tree.bind("<<TreeviewSelect>>", self.load_fuel_data)

    def create_depo_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Yakıt Deposu", padding=5)

        # 1. DEPO DURUM PANELİ
        status_frame = ttk.LabelFrame(frame, text="Anlık Depo Durumu", padding=10)
        status_frame.pack(fill=tk.X, padx=10, pady=5)

        # Yakıt bilgisi etiketi
        self.fuel_level_var = tk.StringVar()
        self.fuel_level_label = ttk.Label(
            status_frame,
            textvariable=self.fuel_level_var,
            font=('Segoe UI', 12, 'bold'),
            foreground="#2c3e50"
        )
        self.fuel_level_label.pack(side=tk.LEFT, padx=10, pady=5)

        # Yakıt seviyesi (cm) etiketi
        self.fuel_cm_var = tk.StringVar()
        self.fuel_cm_label = ttk.Label(
            status_frame,
            textvariable=self.fuel_cm_var,
            font=('Segoe UI', 12, 'bold'),
            foreground="#3498db"
        )
        self.fuel_cm_label.pack(side=tk.LEFT, padx=10, pady=5)

        # Başlangıç değerlerini güncelle
        self.update_fuel_level()

        # 2. DEPO İŞLEMLERİ PANELİ (Sadece giriş işlemi)
        add_frame = ttk.LabelFrame(frame, text="Depo Doldurma İşlemleri", padding=10)
        add_frame.pack(fill=tk.X, padx=10, pady=5)

        # Miktar girişi
        ttk.Label(add_frame, text="Miktar (L):", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.depo_amount_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.depo_amount_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        # Fiyat girişi
        ttk.Label(add_frame, text="Birim Fiyat:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.depo_price_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.depo_price_entry.insert(0, f"{self.current_fuel_price:.2f}")
        self.depo_price_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")

        # Tarih seçici
        ttk.Label(add_frame, text="Tarih:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.depo_date_entry = DateEntry(add_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.depo_date_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")

        # Not alanı
        ttk.Label(add_frame, text="Not:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.depo_notes_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.depo_notes_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")

        # Butonlar
        btn_frame = ttk.Frame(add_frame)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=10)

        ttk.Button(
            btn_frame,
            text="Depoyu Doldur",
            command=self.save_depo_record,
            style='Primary.TButton'
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            btn_frame,
            text="Hesapla",
            command=self.calculate_depo_cost,
            style='TButton'
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            btn_frame,
            text="Sil",
            command=self.delete_depo_record,
            style='Danger.TButton'
        ).pack(side=tk.LEFT, padx=5)

        # 3. DEPO HAREKET LİSTESİ (Sadece giriş işlemleri)
        list_frame = ttk.LabelFrame(frame, text="Depo Doldurma Kayıtları", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        # Treeview sütunları
        columns = ("ID", "Tarih", "Miktar (L)", "Birim Fiyat", "Toplam", "Not")
        self.depo_tree = ttk.Treeview(
            list_frame,
            columns=columns,
            show="headings",
            selectmode="browse"
        )

        # Sütun başlıkları
        for col in columns:
            self.depo_tree.heading(col, text=col, anchor="center")
            self.depo_tree.column(col, width=100, anchor="center")

        # Scrollbar
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.depo_tree.yview)
        self.depo_tree.configure(yscrollcommand=scrollbar.set)

        self.depo_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

        # Satır renkleri
        self.depo_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.depo_tree.tag_configure('evenrow', background="white")

        # Seçim olayı
        self.depo_tree.bind("<<TreeviewSelect>>", self.load_depo_data)

        # Başlangıç verilerini yükle
        self.load_depo_records()

    def create_bakim_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Bakım Kayıtları", padding=5)
        
        # Add maintenance panel
        add_frame = ttk.LabelFrame(frame, text="Bakım Ekle", padding=10)
        add_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Label(add_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_vehicle_combo = ttk.Combobox(add_frame, state="readonly", font=self.normal_font)
        self.maintenance_vehicle_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="KM:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_km_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.maintenance_km_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Tespit Edilen Arıza:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_fault_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.maintenance_fault_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Yapılan İşlem:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_repair_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.maintenance_repair_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="İşçilik Tutarı:", style='Subtitle.TLabel').grid(row=4, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_labor_cost_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.maintenance_labor_cost_entry.grid(row=4, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Malzeme Tutarı:", style='Subtitle.TLabel').grid(row=5, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_material_cost_entry = ttk.Entry(add_frame, font=self.normal_font)
        self.maintenance_material_cost_entry.grid(row=5, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(add_frame, text="Tarih:", style='Subtitle.TLabel').grid(row=6, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_date_entry = DateEntry(add_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.maintenance_date_entry.grid(row=6, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame = ttk.Frame(add_frame)
        btn_frame.grid(row=7, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Kaydet", command=self.save_maintenance, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Temizle", command=self.clear_maintenance_form, style='TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Sil", command=self.delete_maintenance, style='Danger.TButton').pack(side=tk.LEFT, padx=5)
        
        # Maintenance records list
        list_frame = ttk.LabelFrame(frame, text="Bakım Kayıtları", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        columns = ("ID", "Tarih", "Plaka", "KM", "Arıza", "Yapılan İşlem", "İşçilik", "Malzeme", "Toplam")
        self.maintenance_tree = ttk.Treeview(list_frame, columns=columns, show="headings", selectmode="browse")
        
        for col in columns:
            self.maintenance_tree.heading(col, text=col, anchor="center")
            self.maintenance_tree.column(col, width=100, anchor="center")
        
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.maintenance_tree.yview)
        self.maintenance_tree.configure(yscrollcommand=scrollbar.set)
        
        self.maintenance_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.maintenance_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.maintenance_tree.tag_configure('evenrow', background="white")
        
        self.maintenance_tree.bind("<<TreeviewSelect>>", self.load_maintenance_data)

    def create_muayene_bakim_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Muayene ve Bakım", padding=5)
        
        # Main frame
        main_frame = ttk.Frame(frame)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # Left frame - Inspection
        left_frame = ttk.LabelFrame(main_frame, text="Muayene Ekle", padding=10)
        left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        ttk.Label(left_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.inspection_vehicle_combo = ttk.Combobox(left_frame, state="readonly", font=self.normal_font)
        self.inspection_vehicle_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(left_frame, text="KM:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.inspection_km_entry = ttk.Entry(left_frame, font=self.normal_font)
        self.inspection_km_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(left_frame, text="Tarih:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.inspection_date_entry = DateEntry(left_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.inspection_date_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(left_frame, text="Sonraki Muayene:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.next_inspection_date_entry = DateEntry(left_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.next_inspection_date_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame = ttk.Frame(left_frame)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Kaydet", command=self.save_inspection, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Temizle", command=self.clear_inspection_form, style='TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Sil", command=self.delete_inspection, style='Danger.TButton').pack(side=tk.LEFT, padx=5)
        
        # Right frame - Periodic Maintenance
        right_frame = ttk.LabelFrame(main_frame, text="Periyodik Bakım Ekle", padding=10)
        right_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        ttk.Label(right_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_vehicle_combo2 = ttk.Combobox(right_frame, state="readonly", font=self.normal_font)
        self.maintenance_vehicle_combo2.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(right_frame, text="KM:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_km_entry2 = ttk.Entry(right_frame, font=self.normal_font)
        self.maintenance_km_entry2.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(right_frame, text="Tarih:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.maintenance_date_entry2 = DateEntry(right_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.maintenance_date_entry2.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(right_frame, text="Sonraki Bakım KM:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.next_maintenance_km_entry = ttk.Entry(right_frame, font=self.normal_font)
        self.next_maintenance_km_entry.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(right_frame, text="Sonraki Bakım Tarihi:", style='Subtitle.TLabel').grid(row=4, column=0, sticky="w", pady=5, padx=5)
        self.next_maintenance_date_entry = DateEntry(right_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.next_maintenance_date_entry.grid(row=4, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame2 = ttk.Frame(right_frame)
        btn_frame2.grid(row=5, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame2, text="Kaydet", command=self.save_periodic_maintenance, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame2, text="Temizle", command=self.clear_periodic_maintenance_form, style='TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame2, text="Sil", command=self.delete_periodic_maintenance, style='Danger.TButton').pack(side=tk.LEFT, padx=5)
        
        # Inspection records list
        list_frame = ttk.LabelFrame(frame, text="Muayene ve Bakım Kayıtları", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        columns = ("ID", "Tarih", "Plaka", "KM", "Sonraki Muayene", "Sonraki Bakım", "Sonraki Bakım KM")
        self.inspection_tree = ttk.Treeview(list_frame, columns=columns, show="headings", selectmode="browse")
        
        for col in columns:
            self.inspection_tree.heading(col, text=col, anchor="center")
            self.inspection_tree.column(col, width=120, anchor="center")
        
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.inspection_tree.yview)
        self.inspection_tree.configure(yscrollcommand=scrollbar.set)
        
        self.inspection_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.inspection_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.inspection_tree.tag_configure('evenrow', background="white")
        
        self.inspection_tree.bind("<<TreeviewSelect>>", self.load_inspection_data)

    def create_rapor_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Raporlar", padding=5)
        
        # Filter panel
        filter_frame = ttk.LabelFrame(frame, text="Filtrele", padding=10)
        filter_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Label(filter_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.report_vehicle_combo = ttk.Combobox(filter_frame, state="readonly", font=self.normal_font)
        self.report_vehicle_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(filter_frame, text="Başlangıç Tarihi:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.start_date_entry = DateEntry(filter_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.start_date_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(filter_frame, text="Bitiş Tarihi:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.end_date_entry = DateEntry(filter_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.end_date_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(filter_frame, text="Rapor Türü:", style='Subtitle.TLabel').grid(row=3, column=0, sticky="w", pady=5, padx=5)
        self.report_type_combo = ttk.Combobox(filter_frame, values=["Yakıt", "Bakım", "Muayene"], state="readonly", font=self.normal_font)
        self.report_type_combo.grid(row=3, column=1, padx=5, pady=5, sticky="ew")
        self.report_type_combo.current(0)
        
        btn_frame = ttk.Frame(filter_frame)
        btn_frame.grid(row=4, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Filtrele", command=self.filter_reports, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Grafik Oluştur", command=self.show_consumption_graph, style='Success.TButton').pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Excel'e Aktar", command=self.export_filtered_report_to_excel, style='Success.TButton').pack(side=tk.LEFT, padx=5)
        
        # Report list
        list_frame = ttk.LabelFrame(frame, text="Rapor", padding=10)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        self.report_tree = ttk.Treeview(list_frame, show="headings", selectmode="browse")
        self.report_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.report_tree.yview)
        self.report_tree.configure(yscrollcommand=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.report_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.report_tree.tag_configure('evenrow', background="white")

    def create_maliyet_tab(self):
        frame = ttk.Frame(self.notebook)
        self.notebook.add(frame, text="Maliyet Hesaplama", padding=5)
        
        # Cost calculation panel
        cost_frame = ttk.LabelFrame(frame, text="Yakıt Maliyeti Hesaplama", padding=10)
        cost_frame.pack(fill=tk.X, padx=10, pady=5)
        
        ttk.Label(cost_frame, text="Araç:", style='Subtitle.TLabel').grid(row=0, column=0, sticky="w", pady=5, padx=5)
        self.cost_vehicle_combo = ttk.Combobox(cost_frame, state="readonly", font=self.normal_font)
        self.cost_vehicle_combo.grid(row=0, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(cost_frame, text="Başlangıç Tarihi:", style='Subtitle.TLabel').grid(row=1, column=0, sticky="w", pady=5, padx=5)
        self.cost_start_date_entry = DateEntry(cost_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.cost_start_date_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")
        
        ttk.Label(cost_frame, text="Bitiş Tarihi:", style='Subtitle.TLabel').grid(row=2, column=0, sticky="w", pady=5, padx=5)
        self.cost_end_date_entry = DateEntry(cost_frame, date_pattern='dd.mm.yyyy', font=self.normal_font)
        self.cost_end_date_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")
        
        btn_frame = ttk.Frame(cost_frame)
        btn_frame.grid(row=3, column=0, columnspan=2, pady=10)
        ttk.Button(btn_frame, text="Hesapla", command=self.calculate_cost, style='Primary.TButton').pack(side=tk.LEFT, padx=5)
        
        # Results
        result_frame = ttk.LabelFrame(frame, text="Sonuçlar", padding=10)
        result_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        columns = ("Araç Plakası", "Toplam Yakıt (L)", "Toplam Maliyet (TL)", "Ortalama Tüketim (L/100km)")
        self.cost_tree = ttk.Treeview(result_frame, columns=columns, show="headings", selectmode="browse")
        
        for col in columns:
            self.cost_tree.heading(col, text=col, anchor="center")
            self.cost_tree.column(col, width=150, anchor="center")
        
        scrollbar = ttk.Scrollbar(result_frame, orient="vertical", command=self.cost_tree.yview)
        self.cost_tree.configure(yscrollcommand=scrollbar.set)
        
        self.cost_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Alternatif satır renkleri
        self.cost_tree.tag_configure('oddrow', background=self.lighter_bg)
        self.cost_tree.tag_configure('evenrow', background="white")

    def load_initial_data(self):
        self.load_vehicles()
        self.load_fuel_records()
        self.load_depo_records()
        self.load_maintenance_records()
        self.load_inspection_records()
        self.load_reports()
        self.update_vehicle_combos()
        
        self.fuel_price_entry.delete(0, tk.END)
        self.fuel_price_entry.insert(0, f"{self.current_fuel_price:.2f}")

    def update_vehicle_combos(self):
        self.cursor.execute("SELECT plate FROM vehicles ORDER BY plate")
        vehicles = [plate[0] for plate in self.cursor.fetchall()]
        
        combos = [
            'fuel_vehicle_combo', 'maintenance_vehicle_combo', 'inspection_vehicle_combo',
            'maintenance_vehicle_combo2', 'report_vehicle_combo', 'cost_vehicle_combo'
        ]
        
        for combo_name in combos:
            if hasattr(self, combo_name):
                combo = getattr(self, combo_name)
                combo['values'] = vehicles
                if vehicles:
                    combo.current(0)
        
        if hasattr(self, 'report_vehicle_combo'):
            self.report_vehicle_combo['values'] = ["Tüm Araçlar"] + vehicles
            self.report_vehicle_combo.current(0)
            
        if hasattr(self, 'cost_vehicle_combo'):
            self.cost_vehicle_combo['values'] = ["Tüm Araçlar"] + vehicles
            self.cost_vehicle_combo.current(0)

    def load_vehicles(self):
        self.vehicle_tree.delete(*self.vehicle_tree.get_children())
        self.cursor.execute("SELECT plate, model, km, driver FROM vehicles ORDER BY plate")
        for i, row in enumerate(self.cursor.fetchall()):
            tag = 'evenrow' if i % 2 == 0 else 'oddrow'
            self.vehicle_tree.insert("", tk.END, values=row, tags=(tag,))

    def load_fuel_records(self):
        self.fuel_tree.delete(*self.fuel_tree.get_children())
        self.cursor.execute("""
        SELECT f.id, strftime('%d.%m.%Y', f.date), v.plate, f.km, f.amount, f.price, f.total 
        FROM fuel_records f
        JOIN vehicles v ON f.vehicle_id = v.id
        ORDER BY f.date DESC
        LIMIT 100
        """)
        for i, row in enumerate(self.cursor.fetchall()):
            tag = 'evenrow' if i % 2 == 0 else 'oddrow'
            self.fuel_tree.insert("", tk.END, values=row, tags=(tag,))

    def load_depo_records(self):
        """Depo hareketlerini yükler"""
        try:
            self.depo_tree.delete(*self.depo_tree.get_children())
            
            self.cursor.execute("""
            SELECT 
                id,
                strftime('%d.%m.%Y', date),
                amount,
                price,
                total,
                COALESCE(notes, '')
            FROM fuel_tank
            WHERE transaction_type='IN'
            ORDER BY date DESC, id DESC
            LIMIT 200
            """)

            for i, row in enumerate(self.cursor.fetchall()):
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                formatted_row = (
                    row[0],  # ID
                    row[1],  # Tarih
                    f"{row[2]:.2f}",  # Miktar
                    f"{row[3]:.2f}",  # Fiyat
                    f"{row[4]:.2f}",  # Toplam
                    row[5]   # Not
                )
                self.depo_tree.insert("", tk.END, values=formatted_row, tags=(tag,))

            self.update_fuel_level()

        except Exception as e:
            messagebox.showerror("Hata", f"Kayıtlar yüklenirken hata: {str(e)}")

    def load_maintenance_records(self):
        self.maintenance_tree.delete(*self.maintenance_tree.get_children())
        self.cursor.execute("""
        SELECT m.id, strftime('%d.%m.%Y', m.date), v.plate, m.km, m.fault, m.repair, 
               m.labor_cost, m.material_cost, (m.labor_cost + m.material_cost)
        FROM maintenance m
        JOIN vehicles v ON m.vehicle_id = v.id
        ORDER BY m.date DESC
        LIMIT 100
        """)
        for i, row in enumerate(self.cursor.fetchall()):
            tag = 'evenrow' if i % 2 == 0 else 'oddrow'
            self.maintenance_tree.insert("", tk.END, values=row, tags=(tag,))

    def load_inspection_records(self):
        self.inspection_tree.delete(*self.inspection_tree.get_children())
        self.cursor.execute("""
        SELECT i.id, strftime('%d.%m.%Y', i.date), v.plate, i.km, 
               COALESCE(strftime('%d.%m.%Y', i.next_inspection_date), ''), 
               COALESCE(strftime('%d.%m.%Y', i.next_maintenance_date), ''), 
               COALESCE(i.next_maintenance_km, '')
        FROM inspections i
        JOIN vehicles v ON i.vehicle_id = v.id
        ORDER BY i.date DESC
        LIMIT 100
        """)
        for i, row in enumerate(self.cursor.fetchall()):
            tag = 'evenrow' if i % 2 == 0 else 'oddrow'
            self.inspection_tree.insert("", tk.END, values=row, tags=(tag,))

    def load_reports(self):
        self.report_tree.delete(*self.report_tree.get_children())

    def load_vehicle_data(self, event):
        selected = self.vehicle_tree.selection()
        if not selected:
            return
            
        vehicle = self.vehicle_tree.item(selected[0])['values']
        self.plaka_entry.delete(0, tk.END)
        self.plaka_entry.insert(0, vehicle[0])
        self.model_entry.delete(0, tk.END)
        self.model_entry.insert(0, vehicle[1])
        self.km_entry.delete(0, tk.END)
        self.km_entry.insert(0, vehicle[2])
        self.driver_entry.delete(0, tk.END)
        if len(vehicle) > 3:
            self.driver_entry.insert(0, vehicle[3])

    def load_fuel_data(self, event):
        selected = self.fuel_tree.selection()
        if not selected:
            return
            
        fuel = self.fuel_tree.item(selected[0])['values']
        self.fuel_vehicle_combo.set(fuel[2])
        self.fuel_km_entry.delete(0, tk.END)
        self.fuel_km_entry.insert(0, fuel[3])
        self.fuel_amount_entry.delete(0, tk.END)
        self.fuel_amount_entry.insert(0, fuel[4])
        self.fuel_price_entry.delete(0, tk.END)
        self.fuel_price_entry.insert(0, fuel[5])
        self.fuel_date_entry.set_date(datetime.strptime(fuel[1], "%d.%m.%Y"))

    def load_depo_data(self, event):
        selected = self.depo_tree.selection()
        if not selected:
            return

        try:
            record = self.depo_tree.item(selected[0])['values']
            self.depo_amount_entry.delete(0, tk.END)
            self.depo_amount_entry.insert(0, record[2])  # Miktar

            self.depo_price_entry.delete(0, tk.END)
            self.depo_price_entry.insert(0, record[3])  # Birim Fiyat

            self.depo_date_entry.set_date(datetime.strptime(record[1], "%d.%m.%Y"))  # Tarih

            self.depo_notes_entry.delete(0, tk.END)
            if len(record) > 5:  # Not alanı varsa
                self.depo_notes_entry.insert(0, record[5])

        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt yüklenirken hata: {str(e)}")
            self.clear_depo_form()

    def load_maintenance_data(self, event):
        selected = self.maintenance_tree.selection()
        if not selected:
            return
            
        maintenance = self.maintenance_tree.item(selected[0])['values']
        self.maintenance_vehicle_combo.set(maintenance[2])
        self.maintenance_km_entry.delete(0, tk.END)
        self.maintenance_km_entry.insert(0, maintenance[3])
        self.maintenance_fault_entry.delete(0, tk.END)
        self.maintenance_fault_entry.insert(0, maintenance[4])
        self.maintenance_repair_entry.delete(0, tk.END)
        self.maintenance_repair_entry.insert(0, maintenance[5])
        self.maintenance_labor_cost_entry.delete(0, tk.END)
        self.maintenance_labor_cost_entry.insert(0, maintenance[6])
        self.maintenance_material_cost_entry.delete(0, tk.END)
        self.maintenance_material_cost_entry.insert(0, maintenance[7])
        self.maintenance_date_entry.set_date(datetime.strptime(maintenance[1], "%d.%m.%Y"))

    def load_inspection_data(self, event):
        selected = self.inspection_tree.selection()
        if not selected:
            return
            
        inspection = self.inspection_tree.item(selected[0])['values']
        self.inspection_vehicle_combo.set(inspection[2])
        self.inspection_km_entry.delete(0, tk.END)
        self.inspection_km_entry.insert(0, inspection[3])
        self.inspection_date_entry.set_date(datetime.strptime(inspection[1], "%d.%m.%Y"))
        
        # Handle possible None values for dates
        if inspection[4]:  # next_inspection_date
            try:
                self.next_inspection_date_entry.set_date(datetime.strptime(inspection[4], "%d.%m.%Y"))
            except ValueError:
                pass
        
        if inspection[5]:  # next_maintenance_date
            try:
                self.next_maintenance_date_entry.set_date(datetime.strptime(inspection[5], "%d.%m.%Y"))
            except ValueError:
                pass
        
        if inspection[6]:  # next_maintenance_km
            self.next_maintenance_km_entry.delete(0, tk.END)
            self.next_maintenance_km_entry.insert(0, inspection[6])

    def clear_vehicle_form(self):
        self.plaka_entry.delete(0, tk.END)
        self.model_entry.delete(0, tk.END)
        self.km_entry.delete(0, tk.END)
        self.driver_entry.delete(0, tk.END)
        self.vehicle_tree.selection_remove(self.vehicle_tree.selection())

    def clear_maintenance_form(self):
        self.maintenance_km_entry.delete(0, tk.END)
        self.maintenance_fault_entry.delete(0, tk.END)
        self.maintenance_repair_entry.delete(0, tk.END)
        self.maintenance_labor_cost_entry.delete(0, tk.END)
        self.maintenance_material_cost_entry.delete(0, tk.END)
        self.maintenance_tree.selection_remove(self.maintenance_tree.selection())

    def clear_inspection_form(self):
        self.inspection_km_entry.delete(0, tk.END)
        self.inspection_tree.selection_remove(self.inspection_tree.selection())
        next_year = datetime.now() + timedelta(days=365)
        self.next_inspection_date_entry.set_date(next_year)

    def clear_periodic_maintenance_form(self):
        self.maintenance_km_entry2.delete(0, tk.END)
        self.next_maintenance_km_entry.delete(0, tk.END)
        next_year = datetime.now() + timedelta(days=365)
        self.next_maintenance_date_entry.set_date(next_year)

    def clear_depo_form(self):
        self.depo_amount_entry.delete(0, tk.END)
        self.depo_price_entry.delete(0, tk.END)
        self.depo_price_entry.insert(0, f"{self.current_fuel_price:.2f}")
        self.depo_notes_entry.delete(0, tk.END)
        self.transaction_type.set("IN")
        self.depo_tree.selection_remove(self.depo_tree.selection())

    def save_vehicle(self):
        plate = self.plaka_entry.get().strip().upper()
        model = self.model_entry.get().strip()
        km = self.km_entry.get().strip()
        driver = self.driver_entry.get().strip()

        if not plate or not model or not km:
            messagebox.showerror("Hata", "Lütfen zorunlu alanları doldurun (Plaka, Model, KM)!")
            return

        try:
            km = int(km)

            selected = self.vehicle_tree.selection()
            if selected:  # Update
                old_plate = self.vehicle_tree.item(selected[0])['values'][0]
                # Buradaki SQL sorgusunun doğru şekilde girintilendiğinden emin olun
                self.cursor.execute(
                    "UPDATE vehicles SET plate=?, model=?, km=?, driver=? WHERE plate=?",
                    (plate, model, km, driver, old_plate)
                )
                messagebox.showinfo("Başarılı", "Araç bilgileri güncellendi!")
            else:  # Insert
                self.cursor.execute(
                    "INSERT INTO vehicles (plate, model, km, driver) VALUES (?, ?, ?, ?)",
                    (plate, model, km, driver)
                )
                messagebox.showinfo("Başarılı", "Yeni araç eklendi!")

            self.conn.commit()
            self.load_vehicles()
            self.update_vehicle_combos()
            self.clear_vehicle_form()

        except ValueError:
            messagebox.showerror("Hata", "Geçersiz KM değeri!")
        except sqlite3.IntegrityError:
            messagebox.showerror("Hata", "Bu plaka zaten kayıtlı!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()


    def save_fuel_record(self):
        vehicle = self.fuel_vehicle_combo.get()
        km = self.fuel_km_entry.get().strip()
        amount = self.fuel_amount_entry.get().strip()
        price = self.fuel_price_entry.get().strip() or str(self.current_fuel_price)
        date = self.fuel_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format

        if not vehicle or not km or not amount:
            messagebox.showerror("Hata", "Lütfen zorunlu alanları doldurun (Araç, KM, Miktar)!")
            return

        try:
            # Depodaki yakıt miktarını kontrol et
            if float(amount) > self.current_fuel_level:
                messagebox.showerror("Hata", "Depoda yeterli yakıt yok!")
                return

            self.cursor.execute("SELECT id FROM vehicles WHERE plate=?", (vehicle,))
            vehicle_id = self.cursor.fetchone()[0]

            km = int(km)
            amount = float(amount)
            price = float(price)
            total = amount * price

            selected = self.fuel_tree.selection()
            if selected:  # Update
                record_id = self.fuel_tree.item(selected[0])['values'][0]
                self.cursor.execute(
                    "UPDATE fuel_records SET vehicle_id=?, date=?, km=?, amount=?, price=?, total=? WHERE id=?",
                    (vehicle_id, date, km, amount, price, total, record_id))
                messagebox.showinfo("Başarılı", "Yakıt kaydı güncellendi!")
            else:  # Insert
                self.cursor.execute(
                    "INSERT INTO fuel_records (vehicle_id, date, km, amount, price, total) VALUES (?, ?, ?, ?, ?, ?)",
                    (vehicle_id, date, km, amount, price, total))
                messagebox.showinfo("Başarılı", "Yakıt kaydı eklendi!")

                # Depodan yakıt çıkışı yap
                self.cursor.execute("""
                INSERT INTO fuel_tank 
                (date, amount, price, total, transaction_type, notes)
                VALUES (?, ?, ?, ?, ?, ?)
                """, (date, amount, price, total, "OUT", f"{vehicle} aracına yakıt verildi"))

            # Update vehicle km
            self.cursor.execute("UPDATE vehicles SET km=? WHERE id=?", (km, vehicle_id))

            self.conn.commit()
            self.load_fuel_records()
            self.load_depo_records()  # Depo kayıtlarını yenile
            self.load_vehicles()

            # Clear form
            self.fuel_km_entry.delete(0, tk.END)
            self.fuel_amount_entry.delete(0, tk.END)

        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()

    def save_depo_record(self):
        """Depo doldurma işlemini kaydeder - Sadece giriş işlemi"""
        # Verileri al
        amount = self.depo_amount_entry.get().strip()
        price = self.depo_price_entry.get().strip()
        date = self.depo_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format
        notes = self.depo_notes_entry.get().strip()

        # Validasyon
        if not amount or not price:
            messagebox.showerror("Hata", "Lütfen miktar ve fiyat bilgilerini girin!")
            return

        try:
            amount = float(amount)
            price = float(price)

            if amount <= 0:
                messagebox.showerror("Hata", "Miktar 0'dan büyük olmalıdır!")
                return

            if price <= 0:
                messagebox.showerror("Hata", "Fiyat 0'dan büyük olmalıdır!")
                return

            total = amount * price

            # Seçili kaydı kontrol et
            selected = self.depo_tree.selection()

            if selected:  # Kayıt güncelleme
                record_id = self.depo_tree.item(selected[0])['values'][0]
                self.cursor.execute("""
                UPDATE fuel_tank 
                SET date=?, amount=?, price=?, total=?, notes=?
                WHERE id=? AND transaction_type='IN'
                """, (date, amount, price, total, notes, record_id))
                message = "Depo kaydı güncellendi!"
            else:  # Yeni kayıt
                self.cursor.execute("""
                INSERT INTO fuel_tank 
                (date, amount, price, total, transaction_type, notes)
                VALUES (?, ?, ?, ?, 'IN', ?)
                """, (date, amount, price, total, notes))
                message = "Depo doldurma kaydı eklendi!"

            # Yakıt fiyatını güncelle
            self.current_fuel_price = price
            self.fiyat_label.config(text=f"Mevcut Yakıt Fiyatı: {self.current_fuel_price:.2f} TL")
            self.fuel_price_entry.delete(0, tk.END)
            self.fuel_price_entry.insert(0, f"{self.current_fuel_price:.2f}")

            # Değişiklikleri kaydet ve arayüzü güncelle
            self.conn.commit()
            self.load_depo_records()
            self.clear_depo_form()

            messagebox.showinfo("Başarılı", message)

        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()

    def save_maintenance(self):
        vehicle = self.maintenance_vehicle_combo.get()
        km = self.maintenance_km_entry.get().strip()
        fault = self.maintenance_fault_entry.get().strip()
        repair = self.maintenance_repair_entry.get().strip()
        labor_cost = self.maintenance_labor_cost_entry.get().strip() or "0"
        material_cost = self.maintenance_material_cost_entry.get().strip() or "0"
        date = self.maintenance_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format
        
        if not vehicle or not km or not fault or not repair:
            messagebox.showerror("Hata", "Lütfen zorunlu alanları doldurun (Araç, KM, Arıza, İşlem)!")
            return
        
        try:
            self.cursor.execute("SELECT id FROM vehicles WHERE plate=?", (vehicle,))
            vehicle_id = self.cursor.fetchone()[0]
            
            km = int(km)
            labor_cost = float(labor_cost)
            material_cost = float(material_cost)
            
            selected = self.maintenance_tree.selection()
            if selected:  # Update
                record_id = self.maintenance_tree.item(selected[0])['values'][0]
                self.cursor.execute(
                    """UPDATE maintenance 
                    SET vehicle_id=?, date=?, km=?, fault=?, repair=?, labor_cost=?, material_cost=?
                    WHERE id=?""",
                    (vehicle_id, date, km, fault, repair, labor_cost, material_cost, record_id))
                messagebox.showinfo("Başarılı", "Bakım kaydı güncellendi!")
            else:  # Insert
                self.cursor.execute(
                    """INSERT INTO maintenance 
                    (vehicle_id, date, km, fault, repair, labor_cost, material_cost) 
                    VALUES (?, ?, ?, ?, ?, ?, ?)""",
                    (vehicle_id, date, km, fault, repair, labor_cost, material_cost))
                messagebox.showinfo("Başarılı", "Bakım kaydı eklendi!")
            
            # Update vehicle km
            self.cursor.execute("UPDATE vehicles SET km=? WHERE id=?", (km, vehicle_id))
            
            self.conn.commit()
            self.load_maintenance_records()
            self.load_vehicles()
            
            # Clear form
            self.clear_maintenance_form()
            
        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()

    def save_inspection(self):
        vehicle = self.inspection_vehicle_combo.get()
        km = self.inspection_km_entry.get().strip()
        date = self.inspection_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format
        next_date = self.next_inspection_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format
        
        if not vehicle or not km or not date or not next_date:
            messagebox.showerror("Hata", "Lütfen zorunlu alanları doldurun (Araç, KM, Tarih, Sonraki Muayene)!")
            return
        
        try:
            self.cursor.execute("SELECT id FROM vehicles WHERE plate=?", (vehicle,))
            vehicle_id = self.cursor.fetchone()[0]
            
            km = int(km)
            
            selected = self.inspection_tree.selection()
            if selected:  # Update
                record_id = self.inspection_tree.item(selected[0])['values'][0]
                self.cursor.execute(
                    """UPDATE inspections 
                    SET vehicle_id=?, date=?, km=?, next_inspection_date=?
                    WHERE id=?""",
                    (vehicle_id, date, km, next_date, record_id))
                messagebox.showinfo("Başarılı", "Muayene kaydı güncellendi!")
            else:  # Insert
                self.cursor.execute(
                    """INSERT INTO inspections 
                    (vehicle_id, date, km, next_inspection_date) 
                    VALUES (?, ?, ?, ?)""",
                    (vehicle_id, date, km, next_date))
                messagebox.showinfo("Başarılı", "Muayene kaydı eklendi!")
            
            # Update vehicle km
            self.cursor.execute("UPDATE vehicles SET km=? WHERE id=?", (km, vehicle_id))
            
            self.conn.commit()
            self.load_inspection_records()
            self.load_vehicles()
            
            # Clear form
            self.clear_inspection_form()
            
        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()

    def save_periodic_maintenance(self):
        vehicle = self.maintenance_vehicle_combo2.get()
        km = self.maintenance_km_entry2.get().strip()
        date = self.maintenance_date_entry2.get_date().strftime("%Y-%m-%d")  # ISO format
        next_maintenance_km = self.next_maintenance_km_entry.get().strip()
        next_maintenance_date = self.next_maintenance_date_entry.get_date().strftime("%Y-%m-%d")  # ISO format
        
        if not all([vehicle, km, date, next_maintenance_km, next_maintenance_date]):
            messagebox.showerror("Hata", "Lütfen tüm alanları doldurun!")
            return
        
        try:
            self.cursor.execute("SELECT id FROM vehicles WHERE plate=?", (vehicle,))
            vehicle_id = self.cursor.fetchone()[0]
            
            km = int(km)
            next_maintenance_km = int(next_maintenance_km)
            
            # Mevcut kaydı kontrol et
            self.cursor.execute("""
            SELECT id FROM inspections 
            WHERE vehicle_id=? AND date=?
            """, (vehicle_id, date))
            existing_record = self.cursor.fetchone()
            
            if existing_record:  # Güncelleme
                self.cursor.execute("""
                UPDATE inspections 
                SET km=?, next_maintenance_km=?, next_maintenance_date=?
                WHERE id=?
                """, (km, next_maintenance_km, next_maintenance_date, existing_record[0]))
                message = "Periyodik bakım kaydı güncellendi!"
            else:  # Yeni kayıt
                self.cursor.execute("""
                INSERT INTO inspections 
                (vehicle_id, date, km, next_maintenance_km, next_maintenance_date) 
                VALUES (?, ?, ?, ?, ?)
                """, (vehicle_id, date, km, next_maintenance_km, next_maintenance_date))
                message = "Periyodik bakım kaydı eklendi!"
            
            # Aracın KM'sini güncelle
            self.cursor.execute("UPDATE vehicles SET km=? WHERE id=?", (km, vehicle_id))
            
            self.conn.commit()
            self.load_inspection_records()
            self.load_vehicles()
            
            messagebox.showinfo("Başarılı", message)
            self.clear_periodic_maintenance_form()
            
        except ValueError:
            messagebox.showerror("Hata", "KM değerleri sayısal olmalıdır!")
        except Exception as e:
            messagebox.showerror("Hata", f"Kayıt sırasında hata: {str(e)}")
            self.conn.rollback()

    def delete_vehicle(self):
        selected = self.vehicle_tree.selection()
        if not selected:
            messagebox.showerror("Hata", "Lütfen silmek istediğiniz aracı seçin!")
            return
            
        plate = self.vehicle_tree.item(selected[0])['values'][0]
        
        if messagebox.askyesno("Onay", f"{plate} plakalı aracı silmek istediğinize emin misiniz?"):
            try:
                # First delete related records
                self.cursor.execute("DELETE FROM fuel_records WHERE vehicle_id IN (SELECT id FROM vehicles WHERE plate=?)", (plate,))
                self.cursor.execute("DELETE FROM maintenance WHERE vehicle_id IN (SELECT id FROM vehicles WHERE plate=?)", (plate,))
                self.cursor.execute("DELETE FROM inspections WHERE vehicle_id IN (SELECT id FROM vehicles WHERE plate=?)", (plate,))
                
                # Then delete the vehicle
                self.cursor.execute("DELETE FROM vehicles WHERE plate=?", (plate,))
                self.conn.commit()
                
                messagebox.showinfo("Başarılı", "Araç ve ilişkili kayıtlar silindi!")
                self.load_vehicles()
                self.load_fuel_records()
                self.load_maintenance_records()
                self.load_inspection_records()
                self.update_vehicle_combos()
                self.clear_vehicle_form()
                
            except Exception as e:
                messagebox.showerror("Hata", f"Silme işlemi sırasında hata: {str(e)}")
                self.conn.rollback()

    def delete_fuel_record(self):
        selected = self.fuel_tree.selection()
        if not selected:
            messagebox.showerror("Hata", "Lütfen silmek istediğiniz kaydı seçin!")
            return
            
        record_id = self.fuel_tree.item(selected[0])['values'][0]
        
        if messagebox.askyesno("Onay", "Bu yakıt kaydını silmek istediğinize emin misiniz?"):
            try:
                self.cursor.execute("DELETE FROM fuel_records WHERE id=?", (record_id,))
                self.conn.commit()
                messagebox.showinfo("Başarılı", "Yakıt kaydı silindi!")
                self.load_fuel_records()
            except Exception as e:
                messagebox.showerror("Hata", f"Silme işlemi sırasında hata: {str(e)}")
                self.conn.rollback()

    def delete_depo_record(self):
        selected = self.depo_tree.selection()
        if not selected:
            messagebox.showerror("Hata", "Lütfen silmek istediğiniz kaydı seçin!")
            return
            
        record_id = self.depo_tree.item(selected[0])['values'][0]
        
        if not messagebox.askyesno(
            "Onay", 
            f"Bu depo kaydını silmek istediğinize emin misiniz?\n"
            "Bu işlem geri alınamaz!"
        ):
            return
        
        try:
            self.cursor.execute("DELETE FROM fuel_tank WHERE id=?", (record_id,))
            self.conn.commit()
            
            # Arayüzü güncelle
            self.load_depo_records()
            self.clear_depo_form()
            
            messagebox.showinfo("Başarılı", "Depo kaydı silindi!")
        except Exception as e:
            messagebox.showerror("Hata", f"Silme işlemi sırasında hata: {str(e)}")
            self.conn.rollback()

    def delete_maintenance(self):
        selected = self.maintenance_tree.selection()
        if not selected:
            messagebox.showerror("Hata", "Lütfen silmek istediğiniz kaydı seçin!")
            return
            
        record_id = self.maintenance_tree.item(selected[0])['values'][0]
        
        if messagebox.askyesno("Onay", "Bu bakım kaydını silmek istediğinize emin misiniz?"):
            try:
                self.cursor.execute("DELETE FROM maintenance WHERE id=?", (record_id,))
                self.conn.commit()
                messagebox.showinfo("Başarılı", "Bakım kaydı silindi!")
                self.load_maintenance_records()
            except Exception as e:
                messagebox.showerror("Hata", f"Silme işlemi sırasında hata: {str(e)}")
                self.conn.rollback()

    def delete_inspection(self):
        selected = self.inspection_tree.selection()
        if not selected:
            messagebox.showerror("Hata", "Lütfen silmek istediğiniz kaydı seçin!")
            return
            
        record_id = self.inspection_tree.item(selected[0])['values'][0]
        
        if messagebox.askyesno("Onay", "Bu muayene kaydını silmek istediğinize emin misiniz?"):
            try:
                self.cursor.execute("DELETE FROM inspections WHERE id=?", (record_id,))
                self.conn.commit()
                messagebox.showinfo("Başarılı", "Muayene kaydı silindi!")
                self.load_inspection_records()
            except Exception as e:
                messagebox.showerror("Hata", f"Silme işlemi sırasında hata: {str(e)}")
                self.conn.rollback()

    def delete_periodic_maintenance(self):
        self.delete_inspection()

    def calculate_fuel_cost(self):
        try:
            amount = float(self.fuel_amount_entry.get() or 0)
            price = float(self.fuel_price_entry.get() or self.current_fuel_price)
            total = amount * price
            messagebox.showinfo("Hesaplama", f"Toplam maliyet: {total:.2f} TL")
        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")

    def calculate_depo_cost(self):
        try:
            amount = self.depo_amount_entry.get().strip()
            price = self.depo_price_entry.get().strip()
            
            if not amount or not price:
                messagebox.showerror("Hata", "Lütfen miktar ve fiyat bilgilerini girin!")
                return
                
            amount = float(amount)
            price = float(price)
            
            if amount <= 0:
                messagebox.showerror("Hata", "Miktar 0'dan büyük olmalıdır!")
                return
                
            if price <= 0:
                messagebox.showerror("Hata", "Fiyat 0'dan büyük olmalıdır!")
                return
            
            total = amount * price
            
            messagebox.showinfo(
                "Hesaplama Sonucu",
                f"{amount:.2f} litre yakıt giriş işlemi\n"
                f"Birim fiyat: {price:.2f} TL\n"
                f"Toplam maliyet: {total:.2f} TL"
            )
        except ValueError:
            messagebox.showerror("Hata", "Geçersiz sayısal değer!")

    def filter_reports(self):
        vehicle = self.report_vehicle_combo.get()
        start_date = self.start_date_entry.get_date().strftime("%Y-%m-%d") if self.start_date_entry.get() else ""
        end_date = self.end_date_entry.get_date().strftime("%Y-%m-%d") if self.end_date_entry.get() else ""
        report_type = self.report_type_combo.get()
        
        if report_type == "Yakıt":
            self.filter_fuel_reports(vehicle, start_date, end_date)
        elif report_type == "Bakım":
            self.filter_maintenance_reports(vehicle, start_date, end_date)
        elif report_type == "Muayene":
            self.filter_inspection_reports(vehicle, start_date, end_date)

    def filter_fuel_reports(self, vehicle, start_date, end_date):
        query = """
        SELECT strftime('%d.%m.%Y', f.date), v.plate, f.km, f.amount, f.price, f.total 
        FROM fuel_records f
        JOIN vehicles v ON f.vehicle_id = v.id
        WHERE 1=1
        """
        params = []
        
        if vehicle and vehicle != "Tüm Araçlar":
            query += " AND v.plate = ?"
            params.append(vehicle)
        
        if start_date:
            query += " AND f.date >= ?"
            params.append(start_date)
        
        if end_date:
            query += " AND f.date <= ?"
            params.append(end_date)
        
        query += " ORDER BY f.date DESC"
        
        try:
            self.report_tree.delete(*self.report_tree.get_children())
            
            columns = ("Tarih", "Plaka", "KM", "Miktar", "Birim Fiyat", "Toplam")
            self.report_tree['columns'] = columns
            for col in columns:
                self.report_tree.heading(col, text=col)
                self.report_tree.column(col, width=100)
            
            self.cursor.execute(query, params)
            
            for i, row in enumerate(self.cursor.fetchall()):
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                self.report_tree.insert("", tk.END, values=row, tags=(tag,))
                
            messagebox.showinfo("Başarılı", "Yakıt raporu filtrelendi!")
            
        except Exception as e:
            messagebox.showerror("Hata", f"Filtreleme sırasında hata: {str(e)}")

    def filter_maintenance_reports(self, vehicle, start_date, end_date):
        query = """
        SELECT strftime('%d.%m.%Y', m.date), v.plate, m.km, m.fault, m.repair, 
               m.labor_cost, m.material_cost, (m.labor_cost + m.material_cost)
        FROM maintenance m
        JOIN vehicles v ON m.vehicle_id = v.id
        WHERE 1=1
        """
        params = []
        
        if vehicle and vehicle != "Tüm Araçlar":
            query += " AND v.plate = ?"
            params.append(vehicle)
        
        if start_date:
            query += " AND m.date >= ?"
            params.append(start_date)
        
        if end_date:
            query += " AND m.date <= ?"
            params.append(end_date)
        
        query += " ORDER BY m.date DESC"
        
        try:
            self.report_tree.delete(*self.report_tree.get_children())
            
            columns = ("Tarih", "Plaka", "KM", "Arıza", "İşlem", "İşçilik", "Malzeme", "Toplam")
            self.report_tree['columns'] = columns
            for col in columns:
                self.report_tree.heading(col, text=col)
                self.report_tree.column(col, width=100)
            
            self.cursor.execute(query, params)
            
            for i, row in enumerate(self.cursor.fetchall()):
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                self.report_tree.insert("", tk.END, values=row, tags=(tag,))
                
            messagebox.showinfo("Başarılı", "Bakım raporu filtrelendi!")
            
        except Exception as e:
            messagebox.showerror("Hata", f"Filtreleme sırasında hata: {str(e)}")

    def filter_inspection_reports(self, vehicle, start_date, end_date):
        query = """
        SELECT strftime('%d.%m.%Y', i.date), v.plate, i.km, 
               COALESCE(strftime('%d.%m.%Y', i.next_inspection_date), '') as next_inspection,
               COALESCE(strftime('%d.%m.%Y', i.next_maintenance_date), '') as next_maintenance,
               COALESCE(i.next_maintenance_km, '') as next_maintenance_km
        FROM inspections i
        JOIN vehicles v ON i.vehicle_id = v.id
        WHERE 1=1
        """
        params = []
        
        if vehicle and vehicle != "Tüm Araçlar":
            query += " AND v.plate = ?"
            params.append(vehicle)
        
        if start_date:
            query += " AND i.date >= ?"
            params.append(start_date)
        
        if end_date:
            query += " AND i.date <= ?"
            params.append(end_date)
        
        query += " ORDER BY i.date DESC"
        
        try:
            self.report_tree.delete(*self.report_tree.get_children())
            
            columns = ("Tarih", "Plaka", "KM", "Sonraki Muayene", "Sonraki Bakım", "Sonraki Bakım KM")
            self.report_tree['columns'] = columns
            for col in columns:
                self.report_tree.heading(col, text=col)
                self.report_tree.column(col, width=100)
            
            self.cursor.execute(query, params)
            
            for i, row in enumerate(self.cursor.fetchall()):
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                self.report_tree.insert("", tk.END, values=row, tags=(tag,))
                
            messagebox.showinfo("Başarılı", "Muayene raporu filtrelendi!")
            
        except Exception as e:
            messagebox.showerror("Hata", f"Filtreleme sırasında hata: {str(e)}")

    def calculate_cost(self):
        vehicle = self.cost_vehicle_combo.get()
        start_date = self.cost_start_date_entry.get_date().strftime("%Y-%m-%d") if self.cost_start_date_entry.get() else ""
        end_date = self.cost_end_date_entry.get_date().strftime("%Y-%m-%d") if self.cost_end_date_entry.get() else ""

        if not start_date or not end_date:
            messagebox.showerror("Hata", "Lütfen başlangıç ve bitiş tarihlerini seçin!")
            return

        try:
            # Yakıt kayıtlarını tarih sırasına göre al
            query = """
            SELECT 
                v.plate,
                f.amount,
                f.km,
                f.date
            FROM fuel_records f
            JOIN vehicles v ON f.vehicle_id = v.id
            WHERE f.date BETWEEN ? AND ?
            """
            params = [start_date, end_date]

            if vehicle and vehicle != "Tüm Araçlar":
                query += " AND v.plate = ?"
                params.append(vehicle)

            query += " ORDER BY f.date ASC"  # Tarihe göre sırala

            self.cursor.execute(query, params)
            records = self.cursor.fetchall()

            if len(records) < 2:
                messagebox.showinfo("Bilgi", "En az 2 kayıt gereklidir!")
                return

            # Plate'e göre grupla (çoklu araç desteği)
            vehicles = {}
            for record in records:
                plate = record[0]
                if plate not in vehicles:
                    vehicles[plate] = []
                vehicles[plate].append((record[1], record[2]))  # (amount, km)

            self.cost_tree.delete(*self.cost_tree.get_children())

            for plate, fuel_data in vehicles.items():
                total_fuel = 0
                total_km = 0
                avg_consumption = 0

                # Her araç için tüketimi hesapla
                if len(fuel_data) >= 2:
                    total_fuel = sum(amount for amount, _ in fuel_data[1:])  # İlk yakıt alımını hariç tut
                    total_km = fuel_data[-1][1] - fuel_data[0][1]  # Son km - ilk km
                    avg_consumption = (total_fuel / total_km) * 100 if total_km > 0 else 0

                # Toplam maliyet (tüm yakıt alımları dahil)
                total_cost_query = """
                SELECT SUM(f.total) 
                FROM fuel_records f
                JOIN vehicles v ON f.vehicle_id = v.id
                WHERE v.plate = ? AND f.date BETWEEN ? AND ?
                """
                self.cursor.execute(total_cost_query, (plate, start_date, end_date))
                total_cost = self.cursor.fetchone()[0] or 0

                self.cost_tree.insert("", tk.END, values=(
                    plate,
                    f"{sum(amount for amount, _ in fuel_data):.2f}",  # Toplam yakıt
                    f"{total_cost:.2f}",
                    f"{avg_consumption:.2f}" if avg_consumption > 0 else "Hesaplanamadı"
                ))

        except Exception as e:
            messagebox.showerror("Hata", f"Maliyet hesaplanırken hata: {str(e)}")

    def check_notifications(self):
        try:
            today = datetime.now().date()
            
            # Bakım uyarıları (KM bazlı) - Daha hassas kontrol
            self.cursor.execute("""
            SELECT COUNT(*) 
            FROM inspections i
            JOIN vehicles v ON i.vehicle_id = v.id
            WHERE i.next_maintenance_km IS NOT NULL 
            AND v.km >= i.next_maintenance_km - 1000
            """)
            self.maintenance_notification_count = self.cursor.fetchone()[0] or 0
            
            # Muayene uyarıları (tarih bazlı)
            self.cursor.execute("""
            SELECT COUNT(*) 
            FROM inspections i
            WHERE i.next_inspection_date IS NOT NULL
            AND date(i.next_inspection_date) <= date(?, '+60 days')
            """, (today.strftime("%Y-%m-%d"),))
            self.inspection_notification_count = self.cursor.fetchone()[0] or 0
            
            # Menüyü güncelle
            self.update_notification_menu()
            
        except Exception as e:
            print("Uyarı kontrolü sırasında hata:", str(e))
        
        finally:
            # 1 dakika sonra tekrar kontrol et
            self.root.after(60000, self.check_notifications)

    def update_fuel_level(self):
        """Depodaki yakıt seviyesini günceller"""
        try:
            # Toplam yakıt miktarını hesapla
            self.cursor.execute("""
            SELECT SUM(
                CASE 
                    WHEN transaction_type='IN' THEN amount 
                    ELSE -amount 
                END
            ) FROM fuel_tank
            """)
            result = self.cursor.fetchone()
            self.current_fuel_level = float(result[0]) if result and result[0] else 0.0
            
            # Negatif değerleri sıfırla
            if self.current_fuel_level < 0:
                self.current_fuel_level = 0.0
            
            # Arayüzü güncelle
            self.fuel_level_var.set(f"Mevcut Yakıt: {self.current_fuel_level:.2f} litre")
            self.fuel_cm_var.set(f"Seviye: {self.current_fuel_level/22:.2f} cm")
            
        except Exception as e:
            messagebox.showerror("Hata", f"Depo durumu güncellenirken hata: {str(e)}")

    def update_notification_menu(self):
        if self.maintenance_notification_count > 0:
            self.notification_menu.entryconfig(0, 
                label=f"KM Uyarıları ({self.maintenance_notification_count})", 
                foreground='red')
        else:
            self.notification_menu.entryconfig(0, 
                label="KM Uyarıları", 
                foreground="black")
        
        if self.inspection_notification_count > 0:
            self.notification_menu.entryconfig(1, 
                label=f"Tarih Uyarıları ({self.inspection_notification_count})", 
                foreground='red')
        else:
            self.notification_menu.entryconfig(1, 
                label="Tarih Uyarıları", 
                foreground="black")

    def show_maintenance_notifications(self):
        try:
            self.cursor.execute("""
            SELECT v.plate, i.next_maintenance_km, v.km, 
                   (v.km - i.next_maintenance_km) as km_diff,
                   i.next_maintenance_date
            FROM inspections i
            JOIN vehicles v ON i.vehicle_id = v.id
            WHERE i.next_maintenance_km IS NOT NULL 
            ORDER BY km_diff DESC
            """)
            alerts = self.cursor.fetchall()
            
            if not alerts:
                messagebox.showinfo("Bilgi", "Bakım gerektiren araç bulunamadı!")
                return
            
            alert_window = tk.Toplevel(self.root)
            alert_window.title("Bakım Uyarıları (KM Bazlı)")
            alert_window.geometry("800x600")
            
            columns = ("Plaka", "Bakım KM", "Mevcut KM", "KM Farkı", "Sonraki Bakım Tarihi")
            tree = ttk.Treeview(alert_window, columns=columns, show="headings")
            
            for col in columns:
                tree.heading(col, text=col, anchor="center")
                tree.column(col, width=120, anchor="center")
            
            scrollbar = ttk.Scrollbar(alert_window, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=scrollbar.set)
            
            tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            
            for alert in alerts:
                km_diff = alert[3]
                if km_diff >= 0:
                    km_diff_str = f"+{km_diff} (GEÇMİŞ)"
                else:
                    km_diff_str = f"{km_diff} (KALAN)"
                
                next_date = alert[4] if alert[4] else "Belirtilmemiş"
                
                tree.insert("", tk.END, values=(
                    alert[0],  # Plaka
                    alert[1],  # Bakım KM
                    alert[2],  # Mevcut KM
                    km_diff_str,
                    next_date
                ))
            
            # Bilgi etiketi ekleyin
            info_label = ttk.Label(
                alert_window,
                text="Pozitif KM Farkı: Bakım gecikmiş | Negatif KM Farkı: Bakıma kalan KM",
                font=('Segoe UI', 10, 'italic'),
                foreground="#666666"
            )
            info_label.pack(pady=5)
            
            ttk.Button(
                alert_window, 
                text="Kapat", 
                command=alert_window.destroy,
                style='Primary.TButton'
            ).pack(pady=10)
            
        except Exception as e:
            messagebox.showerror("Hata", f"Bakım uyarıları gösterilirken hata oluştu: {str(e)}")
            print("Hata detayı:", str(e))  # Konsola hata detayını yazdır

    def show_inspection_notifications(self):
        today = datetime.now().date()
        sixty_days_later = today + timedelta(days=60)
        
        # Muayene uyarıları
        self.cursor.execute("""
        SELECT v.plate, i.next_inspection_date
        FROM inspections i
        JOIN vehicles v ON i.vehicle_id = v.id
        WHERE i.next_inspection_date IS NOT NULL
        ORDER BY i.next_inspection_date
        """)
        
        inspection_alerts = []
        for alert in self.cursor.fetchall():
            try:
                next_date = datetime.strptime(alert[1], "%Y-%m-%d").date()
                days_diff = (next_date - today).days
                if next_date <= sixty_days_later:
                    inspection_alerts.append((alert[0], alert[1], days_diff))
            except ValueError:
                continue
        
        # Periyodik bakım uyarıları
        self.cursor.execute("""
        SELECT v.plate, i.next_maintenance_date
        FROM inspections i
        JOIN vehicles v ON i.vehicle_id = v.id
        WHERE i.next_maintenance_date IS NOT NULL
        ORDER BY i.next_maintenance_date
        """)
        
        maintenance_alerts = []
        for alert in self.cursor.fetchall():
            try:
                next_date = datetime.strptime(alert[1], "%Y-%m-%d").date()
                days_diff = (next_date - today).days
                if next_date <= sixty_days_later:
                    maintenance_alerts.append((alert[0], alert[1], days_diff))
            except ValueError:
                continue
        
        if not inspection_alerts and not maintenance_alerts:
            messagebox.showinfo("Bilgi", "Yaklaşan muayene veya bakım bulunamadı!")
            return
        
        alert_window = tk.Toplevel(self.root)
        alert_window.title("Muayene ve Bakım Uyarıları")
        alert_window.geometry("800x600")
        
        # Create notebook
        notebook = ttk.Notebook(alert_window)
        
        # Inspection alerts tab
        if inspection_alerts:
            inspection_frame = ttk.Frame(notebook)
            notebook.add(inspection_frame, text="Muayene Uyarıları")
            
            columns = ("Plaka", "Sonraki Muayene", "Kalan Gün")
            tree = ttk.Treeview(inspection_frame, columns=columns, show="headings", selectmode="browse")
            
            for col in columns:
                tree.heading(col, text=col, anchor="center")
                tree.column(col, width=120, anchor="center")
            
            scrollbar = ttk.Scrollbar(inspection_frame, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=scrollbar.set)
            
            tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            
            # Alternatif satır renkleri
            tree.tag_configure('oddrow', background=self.lighter_bg)
            tree.tag_configure('evenrow', background="white")
            
            for i, alert in enumerate(inspection_alerts):
                days_diff = alert[2]
                if days_diff < 0:
                    days_str = f"{abs(days_diff)} gün geçmiş"
                else:
                    days_str = f"{days_diff} gün kaldı"
                
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                tree.insert("", tk.END, values=(
                    alert[0], alert[1], days_str
                ), tags=(tag,))
        
        # Maintenance alerts tab
        if maintenance_alerts:
            maintenance_frame = ttk.Frame(notebook)
            notebook.add(maintenance_frame, text="Bakım Uyarıları")
            
            columns = ("Plaka", "Sonraki Bakım", "Kalan Gün")
            tree = ttk.Treeview(maintenance_frame, columns=columns, show="headings", selectmode="browse")
            
            for col in columns:
                tree.heading(col, text=col, anchor="center")
                tree.column(col, width=120, anchor="center")
            
            scrollbar = ttk.Scrollbar(maintenance_frame, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=scrollbar.set)
            
            tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
            scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
            
            # Alternatif satır renkleri
            tree.tag_configure('oddrow', background=self.lighter_bg)
            tree.tag_configure('evenrow', background="white")
            
            for i, alert in enumerate(maintenance_alerts):
                days_diff = alert[2]
                if days_diff < 0:
                    days_str = f"{abs(days_diff)} gün geçmiş"
                else:
                    days_str = f"{days_diff} gün kaldı"
                
                tag = 'evenrow' if i % 2 == 0 else 'oddrow'
                tree.insert("", tk.END, values=(
                    alert[0], alert[1], days_str
                ), tags=(tag,))
        
        notebook.pack(fill=tk.BOTH, expand=True)
        
        # Close button
        ttk.Button(
            alert_window, 
            text="Kapat", 
            command=alert_window.destroy,
            style='Primary.TButton'
        ).pack(pady=10)

    def show_consumption_graph(self):
        vehicle = self.report_vehicle_combo.get()
        
        if not vehicle or vehicle == "Tüm Araçlar":
            messagebox.showerror("Hata", "Lütfen bir araç seçin!")
            return
            
        try:
            self.cursor.execute("""
            SELECT f.date, f.km, f.amount 
            FROM fuel_records f
            JOIN vehicles v ON f.vehicle_id = v.id
            WHERE v.plate = ?
            ORDER BY f.date
            """, (vehicle,))
            
            data = self.cursor.fetchall()
            
            if len(data) < 2:
                messagebox.showerror("Hata", "Grafik oluşturmak için yeterli veri yok!")
                return
            
            dates = [datetime.strptime(row[0], "%Y-%m-%d") for row in data]
            kms = [row[1] for row in data]
            amounts = [row[2] for row in data]
            
            # Calculate consumption (L/100km)
            consumptions = []
            for i in range(1, len(data)):
                km_diff = kms[i] - kms[i-1]
                if km_diff > 0:
                    consumption = (amounts[i] / km_diff) * 100
                    consumptions.append(consumption)
                else:
                    consumptions.append(0)
            
            # Graph window
            graph_window = tk.Toplevel(self.root)
            graph_window.title(f"{vehicle} - Yakıt Tüketim Grafiği")
            graph_window.geometry("800x600")
            
            fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(8, 6))
            
            # KM graph
            ax1.plot(dates, kms, 'b-', marker='o')
            ax1.set_title('Kilometre Takibi', color=self.dark_text)
            ax1.set_ylabel('KM', color=self.dark_text)
            ax1.tick_params(axis='x', colors=self.dark_text)
            ax1.tick_params(axis='y', colors=self.dark_text)
            ax1.grid(True)
            
            # Consumption graph
            ax2.plot(dates[1:], consumptions, 'r-', marker='o')
            ax2.set_title('Yakıt Tüketimi (L/100km)', color=self.dark_text)
            ax2.set_ylabel('Tüketim', color=self.dark_text)
            ax2.tick_params(axis='x', colors=self.dark_text)
            ax2.tick_params(axis='y', colors=self.dark_text)
            ax2.grid(True)
            
            plt.tight_layout()
            
            # Show graph
            canvas = FigureCanvasTkAgg(fig, master=graph_window)
            canvas.draw()
            canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
            
            # Close button
            ttk.Button(
                graph_window, 
                text="Kapat", 
                command=graph_window.destroy,
                style='Primary.TButton'
            ).pack(pady=10)
            
        except Exception as e:
            messagebox.showerror("Hata", f"Grafik oluşturulurken hata: {str(e)}")

    def show_cost_report(self):
        cost_window = tk.Toplevel(self.root)
        cost_window.title("Toplam Maliyet Raporu")
        cost_window.geometry("800x600")
        
        try:
            # Fuel cost
            self.cursor.execute("SELECT SUM(total) FROM fuel_records")
            fuel_cost = self.cursor.fetchone()[0] or 0
            
            # Maintenance cost
            self.cursor.execute("SELECT SUM(labor_cost + material_cost) FROM maintenance")
            maintenance_cost = self.cursor.fetchone()[0] or 0
            
            # Total cost
            total_cost = fuel_cost + maintenance_cost
            
            # Graph data
            labels = ['Yakıt', 'Bakım']
            sizes = [fuel_cost, maintenance_cost]
            colors = ['#ff9999','#66b3ff']
            
            fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(10, 5))
            
            # Pie chart
            ax1.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%', startangle=90)
            ax1.axis('equal')
            ax1.set_title('Maliyet Dağılımı', color=self.dark_text)
            
            # Bar chart
            ax2.bar(labels, sizes, color=colors)
            ax2.set_title('Maliyetler', color=self.dark_text)
            ax2.set_ylabel('TL', color=self.dark_text)
            ax2.tick_params(axis='x', colors=self.dark_text)
            ax2.tick_params(axis='y', colors=self.dark_text)
            
            plt.tight_layout()
            
            # Show graph
            canvas = FigureCanvasTkAgg(fig, master=cost_window)
            canvas.draw()
            canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
            
            # Total cost info
            ttk.Label(
                cost_window, 
                text=f"Toplam Maliyet: {total_cost:.2f} TL\n"
                     f"Yakıt: {fuel_cost:.2f} TL\n"
                     f"Bakım: {maintenance_cost:.2f} TL",
                font=self.subtitle_font,
                foreground=self.dark_text
            ).pack(pady=10)
            
            # Close button
            ttk.Button(
                cost_window, 
                text="Kapat", 
                command=cost_window.destroy,
                style='Primary.TButton'
            ).pack(pady=10)
            
        except Exception as e:
            messagebox.showerror("Hata", f"Rapor oluşturulurken hata: {str(e)}")
            cost_window.destroy()

    def create_backup(self):
        try:
            backup_dir = filedialog.askdirectory(title="Yedek Kaydedilecek Klasörü Seçin")
            if backup_dir:
                timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                backup_path = os.path.join(backup_dir, f"arac_takip_backup_{timestamp}.db")
                
                # Copy current database
                with open(self.db_path, 'rb') as f:
                    db_content = f.read()
                
                with open(backup_path, 'wb') as f:
                    f.write(db_content)
                
                messagebox.showinfo("Başarılı", f"Yedek başarıyla oluşturuldu:\n{backup_path}")
        except Exception as e:
            messagebox.showerror("Hata", f"Yedek oluşturulurken hata: {str(e)}")

    def restore_backup(self):
        try:
            backup_path = filedialog.askopenfilename(
                title="Yedek Dosyasını Seçin",
                filetypes=[("Veritabanı Dosyaları", "*.db"), ("Tüm Dosyalar", "*.*")]
            )
            
            if backup_path:
                if messagebox.askyesno("Onay", "Yedekten geri yükleme yapılacak. Mevcut veriler silinecek. Devam etmek istiyor musunuz?"):
                    # Close current connection
                    self.conn.close()
                    
                    # Copy backup to main database
                    with open(backup_path, 'rb') as f:
                        backup_content = f.read()
                    
                    with open(self.db_path, 'wb') as f:
                        f.write(backup_content)
                    
                    # Restart application
                    messagebox.showinfo("Başarılı", "Yedek başarıyla geri yüklendi. Uygulama yeniden başlatılacak.")
                    self.root.destroy()
                    
                    # Start new instance
                    root = tk.Tk()
                    app = AraçTakipUygulaması(root)
                    root.mainloop()
                    
        except Exception as e:
            messagebox.showerror("Hata", f"Yedekten geri yükleme sırasında hata: {str(e)}")

    def update_fuel_price(self):
        new_price = simpledialog.askfloat(
            "Yakıt Fiyatı Güncelle", 
            "Yeni yakıt fiyatını girin (TL/L):",
            initialvalue=self.current_fuel_price
        )
        
        if new_price and new_price > 0:
            self.current_fuel_price = new_price
            self.fiyat_label.config(text=f"Mevcut Yakıt Fiyatı: {self.current_fuel_price:.2f} TL")
            self.fuel_price_entry.delete(0, tk.END)
            self.fuel_price_entry.insert(0, f"{self.current_fuel_price:.2f}")
            self.depo_price_entry.delete(0, tk.END)
            self.depo_price_entry.insert(0, f"{self.current_fuel_price:.2f}")

    def export_fuel_to_excel(self):
        try:
            # Verileri al
            self.cursor.execute("""
            SELECT strftime('%d.%m.%Y', f.date), v.plate, f.km, f.amount, f.price, f.total 
            FROM fuel_records f
            JOIN vehicles v ON f.vehicle_id = v.id
            ORDER BY f.date DESC
            """)
            fuel_data = self.cursor.fetchall()
            
            if not fuel_data:
                messagebox.showwarning("Uyarı", "Dışa aktarılacak yakıt kaydı bulunamadı!")
                return
                
            # Excel dosyası oluştur
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Yakıt İşlemleri"
            
            # Başlıklar
            headers = ["Tarih", "Plaka", "KM", "Miktar (L)", "Birim Fiyat (TL)", "Toplam (TL)"]
            ws.append(headers)
            
            # Verileri ekle
            for row in fuel_data:
                ws.append(row)
            
            # Stil ayarları
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), 
                                right=Side(style='thin'), 
                                top=Side(style='thin'), 
                                bottom=Side(style='thin'))
            
            # Başlık stilini ayarla
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
            
            # Sütun genişliklerini ayarla
            column_widths = [15, 15, 15, 15, 15, 15]
            for i, column_width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = column_width
            
            # Sayı formatları
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=4, max_col=6):
                for cell in row:
                    cell.number_format = '0.00'
            
            # Dosyayı kaydet
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel Dosyaları", "*.xlsx"), ("Tüm Dosyalar", "*.*")],
                title="Yakıt İşlemlerini Kaydet"
            )
            
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Başarılı", f"Yakıt işlemleri başarıyla Excel'e aktarıldı:\n{file_path}")
                
        except Exception as e:
            messagebox.showerror("Hata", f"Excel'e aktarım sırasında hata: {str(e)}")

    def export_maintenance_to_excel(self):
        try:
            # Verileri al
            self.cursor.execute("""
            SELECT strftime('%d.%m.%Y', m.date), v.plate, m.km, m.fault, m.repair, 
                   m.labor_cost, m.material_cost, (m.labor_cost + m.material_cost)
            FROM maintenance m
            JOIN vehicles v ON m.vehicle_id = v.id
            ORDER BY m.date DESC
            """)
            maintenance_data = self.cursor.fetchall()
            
            if not maintenance_data:
                messagebox.showwarning("Uyarı", "Dışa aktarılacak bakım kaydı bulunamadı!")
                return
                
            # Excel dosyası oluştur
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Bakım İşlemleri"
            
            # Başlıklar
            headers = ["Tarih", "Plaka", "KM", "Arıza", "Yapılan İşlem", 
                      "İşçilik Tutarı (TL)", "Malzeme Tutarı (TL)", "Toplam Tutar (TL)"]
            ws.append(headers)
            
            # Verileri ekle
            for row in maintenance_data:
                ws.append(row)
            
            # Stil ayarları
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), 
                                right=Side(style='thin'), 
                                top=Side(style='thin'), 
                                bottom=Side(style='thin'))
            
            # Başlık stilini ayarla
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
            
            # Sütun genişliklerini ayarla
            column_widths = [15, 15, 15, 30, 30, 15, 15, 15]
            for i, column_width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = column_width
            
            # Sayı formatları
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=6, max_col=8):
                for cell in row:
                    cell.number_format = '0.00'
            
            # Dosyayı kaydet
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel Dosyaları", "*.xlsx"), ("Tüm Dosyalar", "*.*")],
                title="Bakım İşlemlerini Kaydet"
            )
            
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Başarılı", f"Bakım işlemleri başarıyla Excel'e aktarıldı:\n{file_path}")
                
        except Exception as e:
            messagebox.showerror("Hata", f"Excel'e aktarım sırasında hata: {str(e)}")

    def export_cost_report_to_excel(self):
        try:
            # Verileri al (son maliyet raporundaki verileri kullan)
            items = []
            for item in self.cost_tree.get_children():
                items.append(self.cost_tree.item(item)['values'])
            
            if not items:
                messagebox.showwarning("Uyarı", "Dışa aktarılacak maliyet raporu bulunamadı!")
                return
                
            # Excel dosyası oluştur
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Maliyet Raporu"
            
            # Başlıklar
            headers = ["Araç Plakası", "Toplam Yakıt (L)", "Toplam Maliyet (TL)", "Ortalama Tüketim (L/100km)"]
            ws.append(headers)
            
            # Verileri ekle
            for row in items:
                ws.append(row)
            
            # Stil ayarları
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), 
                                right=Side(style='thin'), 
                                top=Side(style='thin'), 
                                bottom=Side(style='thin'))
            
            # Başlık stilini ayarla
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
            
            # Sütun genişliklerini ayarla
            column_widths = [20, 20, 20, 20]
            for i, column_width in enumerate(column_widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = column_width
            
            # Sayı formatları
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=4):
                for cell in row:
                    cell.number_format = '0.00'
            
            # Toplam satırı ekle
            if len(items) > 1:
                ws.append([])  # Boş satır
                total_row = ["TOPLAM"]
                
                # Toplam yakıt
                total_fuel = sum(float(item[1]) for item in items)
                total_row.append(total_fuel)
                
                # Toplam maliyet
                total_cost = sum(float(item[2]) for item in items)
                total_row.append(total_cost)
                
                # Ortalama tüketim (ağırlıklı ortalama)
                total_row.append("")  # Bu hesaplama daha karmaşık olabilir
                
                ws.append(total_row)
                
                # Toplam satırı stilini ayarla
                for col in range(1, len(total_row)+1):
                    cell = ws.cell(row=ws.max_row, column=col)
                    cell.font = Font(bold=True)
                    if col > 1 and col < 4:
                        cell.number_format = '0.00'
            
            # Dosyayı kaydet
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel Dosyaları", "*.xlsx"), ("Tüm Dosyalar", "*.*")],
                title="Maliyet Raporunu Kaydet"
            )
            
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Başarılı", f"Maliyet raporu başarıyla Excel'e aktarıldı:\n{file_path}")
                
        except Exception as e:
            messagebox.showerror("Hata", f"Excel'e aktarım sırasında hata: {str(e)}")
            
    def export_filtered_report_to_excel(self):
        try:
            # Mevcut rapor türünü al
            report_type = self.report_type_combo.get()
            
            if report_type == "Yakıt":
                sheet_name = "Yakıt İşlemleri"
                headers = ["Tarih", "Plaka", "KM", "Miktar (L)", "Birim Fiyat (TL)", "Toplam (TL)"]
                data = []
                for item in self.report_tree.get_children():
                    data.append(self.report_tree.item(item)['values'])
            
            elif report_type == "Bakım":
                sheet_name = "Bakım İşlemleri"
                headers = ["Tarih", "Plaka", "KM", "Arıza", "İşlem", 
                          "İşçilik Tutarı (TL)", "Malzeme Tutarı (TL)", "Toplam Tutar (TL)"]
                data = []
                for item in self.report_tree.get_children():
                    data.append(self.report_tree.item(item)['values'])
            
            elif report_type == "Muayene":
                sheet_name = "Muayene Kayıtları"
                headers = ["Tarih", "Plaka", "KM", "Sonraki Muayene", "Sonraki Bakım", "Sonraki Bakım KM"]
                data = []
                for item in self.report_tree.get_children():
                    data.append(self.report_tree.item(item)['values'])
            
            else:
                messagebox.showwarning("Uyarı", "Geçerli bir rapor türü seçin!")
                return
                
            if not data:
                messagebox.showwarning("Uyarı", "Dışa aktarılacak veri bulunamadı!")
                return
                
            # Excel dosyası oluştur
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            # Başlıkları ekle
            ws.append(headers)
            
            # Verileri ekle
            for row in data:
                ws.append(row)
            
            # Stil ayarları
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="2c3e50", end_color="2c3e50", fill_type="solid")
            thin_border = Border(left=Side(style='thin'), 
                                right=Side(style='thin'), 
                                top=Side(style='thin'), 
                                bottom=Side(style='thin'))
            
            # Başlık stilini ayarla
            for col in range(1, len(headers)+1):
                cell = ws.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
            
            # Sütun genişliklerini ayarla
            for col in range(1, len(headers)+1):
                max_length = max(
                    len(str(headers[col-1])) if headers[col-1] else 0,
                    max(len(str(row[col-1])) for row in data) if data else 0
                )
                ws.column_dimensions[get_column_letter(col)].width = min(max_length + 2, 30)
            
            # Sayı formatları
            if report_type in ["Yakıt", "Bakım"]:
                num_cols = range(3, len(headers)) if report_type == "Yakıt" else range(5, len(headers))
                for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                    for col in num_cols:
                        try:
                            float(row[col].value)
                            row[col].number_format = '0.00'
                        except (ValueError, TypeError):
                            pass
            
            # Dosyayı kaydet
            file_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel Dosyaları", "*.xlsx"), ("Tüm Dosyalar", "*.*")],
                title=f"{report_type} Raporunu Kaydet"
            )
            
            if file_path:
                wb.save(file_path)
                messagebox.showinfo("Başarılı", f"{report_type} raporu başarıyla Excel'e aktarıldı:\n{file_path}")
                
        except Exception as e:
            messagebox.showerror("Hata", f"Excel'e aktarım sırasında hata: {str(e)}")

    def on_closing(self):
        if messagebox.askokcancel("Çıkış", "Uygulamadan çıkmak istediğinize emin misiniz?"):
            try:
                if hasattr(self, 'conn'):
                    self.conn.close()
            except:
                pass
            finally:
                self.root.destroy()


if __name__ == "__main__":
    try:
        root = tk.Tk()
        app = AraçTakipUygulaması(root)
        root.protocol("WM_DELETE_WINDOW", app.on_closing)
        root.mainloop()
    except Exception as e:
        messagebox.showerror("Kritik Hata", f"Uygulama başlatılamadı: {str(e)}")
